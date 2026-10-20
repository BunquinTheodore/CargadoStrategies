"""Build the Word version of JUDGEMENT-REPORT.md.

Usage:  python make_report.py [source.md] [output.docx]

Reads the Markdown report (headings, paragraphs, bullets, tables, images,
**bold**, *italic*, [links](url)) and writes a landscape Word document:
gold #B8913A headings, dark grey table headers with white text, no blue
anywhere (hyperlinks are gold-dark and underlined), images scaled to fit.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
DEFAULT_SRC = HERE / "JUDGEMENT-REPORT.md"
DEFAULT_OUT = Path(r"C:\Users\THEODORE VON JOSHUA\Downloads\Cargado-Strategies-Website-Judgement.docx")

GOLD = RGBColor(0xB8, 0x91, 0x3A)
GOLD_DARK = RGBColor(0x7A, 0x5C, 0x12)
INK = RGBColor(0x1A, 0x1A, 0x1C)
HEADER_FILL = "333333"
ZEBRA_FILL = "F5F0E6"
BORDER = "BFB6A3"

DISPLAY_FONT = "Times New Roman"
BODY_FONT = "Calibri"

PAGE_W_IN, PAGE_H_IN, MARGIN_IN = 11.0, 8.5, 0.7
CONTENT_W_IN = PAGE_W_IN - 2 * MARGIN_IN
MAX_IMG_W_IN = 6.8
MAX_IMG_H_IN = 5.0

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
IMAGE_RE = re.compile(r"^!\[(.*)\]\(([^()\s]+)\)\s*$")


# ---------------------------------------------------------------- styles
def set_font(run_or_style_font, name, size=None, color=None, bold=None, italic=None):
    f = run_or_style_font
    f.name = name
    if size is not None:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = color
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def east_asian_font(style, name):
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_font(normal.font, BODY_FONT, 10.5, INK)
    east_asian_font(normal, BODY_FONT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    sizes = {1: 22, 2: 16, 3: 13}
    for level, size in sizes.items():
        st = doc.styles[f"Heading {level}"]
        set_font(st.font, DISPLAY_FONT, size, GOLD, bold=(level != 1), italic=False)
        east_asian_font(st, DISPLAY_FONT)
        st.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        st = doc.styles[name]
        set_font(st.font, BODY_FONT, 10.5, INK)
        st.paragraph_format.space_after = Pt(3)

    # Hyperlink character style without the theme blue.
    try:
        link = doc.styles["Hyperlink"]
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE
        link = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    link.font.color.rgb = GOLD_DARK
    link.font.underline = True


def configure_page(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(PAGE_W_IN), Inches(PAGE_H_IN)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(MARGIN_IN))
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Cargado Strategies · Website Judgement Report · Round 5 · page ")
    set_font(run.font, BODY_FONT, 8.5, GOLD_DARK)
    add_page_field(footer)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_font(run.font, BODY_FONT, 8.5, GOLD_DARK)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


# ---------------------------------------------------------------- inline text
def add_hyperlink(paragraph, text, url, bold=False, size=None):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "7A5C12")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, bold=False, color=None, size=None):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        link = LINK_RE.fullmatch(part)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2), bold=bold, size=size)
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_italic = not is_bold and part.startswith("*") and part.endswith("*") and len(part) > 2
        content = part[2:-2] if is_bold else part[1:-1] if is_italic else part
        if is_bold and LINK_RE.search(content):
            add_inline(paragraph, content, bold=True, color=color, size=size)
            continue
        run = paragraph.add_run(content)
        run.bold = bold or is_bold
        run.italic = is_italic or None
        if color is not None:
            run.font.color.rgb = color
        if size is not None:
            run.font.size = Pt(size)


# ---------------------------------------------------------------- blocks
def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tblpr.append(borders)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def column_widths(rows, n_cols):
    """Share the content width by the longest text in each column (bounded)."""
    lengths = []
    for c in range(n_cols):
        longest = max(len(re.sub(r"[*\[\]()]", "", r[c])) if c < len(r) else 0 for r in rows)
        header_len = len(rows[0][c]) + 3 if c < len(rows[0]) else 0
        lengths.append(min(max(longest, header_len, 11), 90))
    total = sum(lengths)
    widths = [CONTENT_W_IN * l / total for l in lengths]
    min_w = 0.9
    short = [i for i, w in enumerate(widths) if w < min_w]
    if short:
        spare = CONTENT_W_IN - min_w * len(short)
        rest = sum(w for i, w in enumerate(widths) if i not in short)
        widths = [min_w if i in short else w * spare / rest for i, w in enumerate(widths)]
    return widths


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    n = len(header)
    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = column_widths(rows, n)
    small = n >= 6 or any(len(c) > 180 for r in body for c in r)
    size = 9 if small else 9.5

    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        shade(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        add_inline(p, text.replace("**", ""), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=size)
    repeat_header(table.rows[0])

    for r_i, row in enumerate(body):
        cells = table.add_row().cells
        for i in range(n):
            text = row[i] if i < len(row) else ""
            p = cells[i].paragraphs[0]
            add_inline(p, text, size=size)
            if r_i % 2 == 1:
                shade(cells[i], ZEBRA_FILL)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, caption, rel_path, base):
    path = (base / rel_path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Image not found: {path}")
    with Image.open(path) as im:
        w_px, h_px = im.size
    width = MAX_IMG_W_IN
    height = width * h_px / w_px
    if height > MAX_IMG_H_IN:
        height = MAX_IMG_H_IN
        width = height * w_px / h_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"Screenshot: {caption}")
    set_font(run.font, BODY_FONT, 9, GOLD_DARK, italic=True)
    return path


def add_title_block(doc, title, subtitle_lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    set_font(run.font, DISPLAY_FONT, 30, GOLD)
    rule = doc.add_paragraph()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "B8913A")):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    rule._p.get_or_add_pPr().append(pbdr)
    for line in subtitle_lines:
        sp = doc.add_paragraph()
        add_inline(sp, line, color=GOLD_DARK if line is subtitle_lines[0] else None)


# ---------------------------------------------------------------- parser
def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build(src, out):
    base = src.parent
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    images, tables = [], 0

    title = lines[0].lstrip("# ").strip()
    i = 1
    intro = []
    while i < len(lines) and not lines[i].startswith("## "):
        if lines[i].strip():
            intro.append(lines[i].strip())
        i += 1
    add_title_block(doc, title, intro)

    para_buf = []

    def flush():
        if para_buf:
            add_inline(doc.add_paragraph(), " ".join(para_buf))
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush()
            i += 1
            continue
        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(stripped.lstrip("#").strip(), level=min(level - 1, 3))
            i += 1
            continue
        if stripped.startswith("|"):
            flush()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            tables += 1
            continue
        img = IMAGE_RE.match(stripped)
        if img:
            flush()
            images.append(add_image(doc, img.group(1), img.group(2), base))
            i += 1
            continue
        bullet = re.match(r"^(\s*)[-*] (.*)$", line)
        if bullet:
            flush()
            indent = len(bullet.group(1))
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            add_inline(doc.add_paragraph(style=style), bullet.group(2))
            i += 1
            continue
        numbered = re.match(r"^\s*\d+\. (.*)$", line)
        if numbered:
            flush()
            add_inline(doc.add_paragraph(style="List Number"), numbered.group(1))
            i += 1
            continue
        para_buf.append(stripped)
        i += 1
    flush()

    doc.core_properties.title = title
    doc.core_properties.subject = "Round 5 judgement of the Cargado Strategies website"
    doc.core_properties.author = "Cargado Strategies website review panel"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return images, tables


def verify(out, expected_images, expected_tables):
    doc = Document(str(out))
    pics = len(doc.inline_shapes)
    tables = len(doc.tables)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    xml = doc.part.blob.decode("utf-8", "ignore") if hasattr(doc.part, "blob") else ""
    assert pics == expected_images, f"expected {expected_images} images, found {pics}"
    assert tables == expected_tables, f"expected {expected_tables} tables, found {tables}"
    for blue in ("0000FF", "0563C1", "365F91", "4F81BD", "1F497D", "2E74B5"):
        assert blue not in xml.upper(), f"blue colour {blue} found in document body"
    return pics, tables, len(headings)


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    images, tables = build(src, out)
    pics, n_tables, n_headings = verify(out, len(images), tables)
    print(f"OK {out}  images={pics} tables={n_tables} headings={n_headings} bytes={out.stat().st_size}")


if __name__ == "__main__":
    main()
